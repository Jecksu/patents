from __future__ import annotations

import tempfile
import unittest
import warnings
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document

from scripts import generate_technical_disclosure_docs as generator


EXPECTED_TITLE = (
    "一种面向跨CAD内核的基于轻量摘要闸门和受限局部BREP验证的"
    "CAD模型差异对比方法及系统"
)
EXPECTED_OUTPUT_NAME = (
    "跨CAD内核轻量化摘要-CAD模型差异快速筛选与对比-"
    "专利技术交底书-强化版.docx"
)
EXPECTED_FIGURES = [
    ("lightweight_strengthened-fig1-method.png", "图1 方法闭环流程示意图"),
    ("lightweight_strengthened-fig2-system.png", "图2 系统模块与比较会话数据流示意图"),
    ("lightweight_strengthened-fig3-decision.png", "图3 多变换假设与摘要闸门示意图"),
    (
        "lightweight_strengthened-fig4-candidate-closure.png",
        "图4 候选关系消歧、局部验证和差异闭包示意图",
    ),
    ("lightweight_strengthened-fig5-scenarios.png", "图5 场景处理与降级示意图"),
]


def section_text(section: str) -> str:
    return "\n".join(generator.LIGHTWEIGHT[section])


def docx_text(path: Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(chunks)


class LightweightDisclosureConfigurationTests(unittest.TestCase):
    def test_strengthened_title_output_and_five_figures(self) -> None:
        material = generator.LIGHTWEIGHT
        self.assertEqual(material["full_title"], EXPECTED_TITLE)
        self.assertEqual(material["output"].name, EXPECTED_OUTPUT_NAME)
        self.assertTrue(material.get("require_all_figures", False))
        self.assertEqual(
            [(path.name, caption) for path, caption in material["figures"]],
            EXPECTED_FIGURES,
        )

    def test_background_states_problem_and_prior_art_boundary(self) -> None:
        text = section_text("background")
        required = [
            "跨CAD内核",
            "永久命名",
            "实体笛卡尔积",
            "拓扑重构",
            "位姿",
            "单位",
            "镜像",
            "仿射",
            "对称重复实体",
            "设计语义",
            "CN119312543A",
            "可视化不一致",
            "可视化一致",
            "完整/分层BREP",
            "实体候选",
            "稀疏候选图",
            "互斥消歧",
            "候选邻域BREP",
            "保留待复核端点的差异闭包",
            "全局差异不默认终止局部定位",
        ]
        for term in required:
            with self.subTest(term=term):
                self.assertIn(term, text)
        for internal_wording in [
            "本地已知",
            "规避",
            "专利通过",
            "review-aware",
            "不以“粗比较后精比较”作为核心",
        ]:
            self.assertNotIn(internal_wording, text)
        self.assertIn(
            "CN119312543A公开的流程包括：可视化不一致时判定不同，"
            "可视化一致时进入完整/分层BREP比较",
            text,
        )
        self.assertIn(
            "本方案由摘要闸门形成实体候选和排除关系",
            text,
        )

    def test_solution_covers_ordered_gate_matching_and_local_brep_process(self) -> None:
        text = section_text("solution")
        required = [
            "U_A^l",
            "U_B^l",
            "瞬时定位符",
            "不称为永久命名",
            "初始阶段不执行两个完整模型BREP实体逐项比较",
            "G_X",
            "L_X(e)",
            "SEM_X(e)",
            "H_X(e)",
            "刚体",
            "相似变换",
            "仿射",
            "镜像",
            "对称多变换假设",
            "至少两个顺序摘要闸门",
            "C_gate",
            "最佳/次佳裕量",
            "双向排名",
            "端点互斥",
            "局部最大权匹配",
            "至多h阶邻域",
            "实体数预算",
            "字节数预算",
            "查询次数预算",
            "邻域阶数预算",
            "REVIEW端点",
            "面局部验证",
            "边局部验证",
            "体壳局部验证",
        ]
        for term in required:
            with self.subTest(term=term):
                self.assertIn(term, text)

    def test_solution_defines_lineage_closure_and_change_classification(self) -> None:
        text = section_text("solution")
        required = [
            "R_11",
            "R_1m⊆U_A^l×P_+(U_B^l)",
            "R_m1⊆P_+(U_A^l)×U_B^l",
            "组基数至少为2",
            "三类关系两侧端点分别互斥",
            "旧修剪面因新增拓扑分界对应多个新修剪面",
            "分裂时，新模型相对于旧模型新增内部分界",
            "合并时，旧模型中的内部分界在新模型中删除",
            "存在共同参数化时，以参数域内各新修剪域的并集覆盖原修剪域",
            "不存在共同参数化时，将各修剪域映射到三维后验证裁剪曲面区域覆盖",
            "存在共同参数化时，以参数域内各旧修剪域的并集覆盖新修剪域",
            "将各旧修剪域映射到三维后验证其裁剪曲面区域覆盖新修剪面",
            "度量守恒",
            "支承几何",
            "内部分界",
            "外部邻接",
            "重复键不是分裂/合并",
            "周期缝边",
            "等价参数化",
            "R=R_11∪R_1m∪R_m1",
            "COMMON_11=R_11",
            "TRACK_A=dom(R)",
            "TRACK_B=ran(R)",
            "REVIEW_A",
            "REVIEW_B",
            "REVIEW_A⊆U_A^l",
            "REVIEW_B⊆U_B^l",
            "DEL=U_A^l-(dom(R)∪REVIEW_A)",
            "ADD=U_B^l-(ran(R)∪REVIEW_B)",
            "R_11不等于未修改",
            "tau_same",
            "主分类端点互斥",
            "双侧定位",
            "装配重复发生",
            "顺序重排",
            "证据记录",
        ]
        for term in required:
            with self.subTest(term=term):
                self.assertIn(term, text)
        for inaccurate_wording in ["不虚构", "可消去", "反向条件"]:
            self.assertNotIn(inaccurate_wording, text)

    def test_solution_contains_required_formulas_and_missing_data_rules(self) -> None:
        text = section_text("solution")
        required = [
            "d_num(x,y)=clip((|x-y|-eps_abs)/(eps_abs+eps_rel*max(|x|,|y|)),0,1)",
            "分母为0时采用该维度最小正尺度",
            "I_ms=Σ_t min(c_X(t),c_Y(t))",
            "U_ms=Σ_t max(c_X(t),c_Y(t))",
            "d_ms=1-I_ms/max(U_ms,1)",
            "不得以普通集合Jaccard替代多重集",
            "S(a,b)=Σ alpha_i(1-d_i)/Σ alpha_i",
            "D(a,b)=Σ w_i d_i/Σ w_i",
            "J_cand(a,b)",
            "J_verify(a,b)",
            "alpha_i≥0",
            "Σ_{i∈J_cand(a,b)}alpha_i>0",
            "w_i≥0",
            "Σ_{i∈J_verify(a,b)}w_i>0",
            "候选相似度",
            "验证后差异",
            "缺失维度从分子和分母剔除并重新归一化",
            "全缺失进入低信息/REVIEW",
            "硬冲突直接排除",
            "rho_pair=|C_gate|/(|U_A^l|·|U_B^l|)",
            "候选保留率",
            "1-rho_pair",
            "rho_brep=B_local/(B_full,A+B_full,B)",
            "B_local为本次比较实际读取的局部BREP数据量",
            "三者采用同一计量单位",
            "元数据、离线标定、历史基准或采样估计",
            "分母为0时指标不适用",
        ]
        for term in required:
            with self.subTest(term=term):
                self.assertIn(term, text)

    def test_solution_contains_recomputable_quantitative_example(self) -> None:
        text = section_text("solution")
        required = [
            "示例参数仅用于复算，不构成对实际性能限定",
            "100",
            "105",
            "10,500",
            "310",
            "126",
            "rho_pair=1.2%",
            "94个R_11",
            "1个分裂关系消耗模型A侧1个、模型B侧2个端点",
            "1个合并关系消耗模型A侧2个、模型B侧1个端点",
            "两侧均为97个已跟踪端点",
            "REVIEW为空",
            "DEL=3",
            "ADD=8",
            "86个保持不变",
            "8个修改",
            "80MiB",
            "88MiB",
            "B_local=8.4MiB",
            "rho_brep=5%",
            "用于说明指标计算口径",
        ]
        for term in required:
            with self.subTest(term=term):
                self.assertIn(term, text)
        self.assertNotIn("实测承诺", text)

    def test_solution_covers_scenario_degradation_matrix(self) -> None:
        text = section_text("solution")
        required = [
            "刚体场景",
            "单位变化",
            "未知仿射",
            "对称/主轴",
            "镜像场景",
            "修剪面分裂",
            "修剪面合并/几何修复",
            "周期曲面缝边",
            "解析曲面与NURBS等价",
            "孔/倒角/圆角",
            "装配重复发生与重排",
            "摘要缺失",
            "BREP读取失败/预算超限",
            "非流形/无效拓扑",
            "整体不相关",
            "低置信REVIEW",
        ]
        for term in required:
            with self.subTest(term=term):
                self.assertIn(term, text)

    def test_advantages_and_key_points_capture_combined_invention(self) -> None:
        advantages = section_text("advantages")
        for term in [
            "候选对压缩",
            "按需局部读取",
            "跨内核拓扑谱系",
            "低信息不强制定性",
            "实体级和语义级双层定位",
            "可审计指标",
        ]:
            with self.subTest(section="advantages", term=term):
                self.assertIn(term, advantages)

        key_points = section_text("key_points")
        for term in [
            "实体候选关系和排除关系",
            "稀疏候选图",
            "双向裕量",
            "端点互斥",
            "判定需要驱动",
            "预算约束的局部BREP",
            "R_11、R_1m和R_m1",
            "REVIEW保护闭包",
            "语义聚合",
            "双侧定位",
            "rho_pair和rho_brep",
        ]:
            with self.subTest(section="key_points", term=term):
                self.assertIn(term, key_points)


class PermanentDisclosureRegressionTests(unittest.TestCase):
    def test_permanent_material_retains_current_user_changes(self) -> None:
        material = generator.PERMANENT
        self.assertEqual(
            material["output"].name,
            "永久命名-CAD模型差异对比-专利技术交底书-权利要求增强版.docx",
        )
        self.assertEqual(
            material["figures"][3],
            (
                generator.ASSET_DIR / "permanent_naming-fig4-metrics-diff.png",
                "图4 命名域指标计算与模型具体差异计算流程图",
            ),
        )
        permanent_text = "\n".join(
            paragraph
            for section in ("background", "solution", "advantages", "key_points")
            for paragraph in material[section]
        )
        for term in [
            "C_pair=2|K_A∩K_B|/(|K_A|+|K_B|)",
            "Q_cons=Σ(q_k·1[q_k≥τ_cons])/|K_A∩K_B|",
            "C_conf=w_1C_intra+w_2C_cross+w_3C_map",
            "模型A有100个可审计实体",
            "C_cross约为3.26%",
        ]:
            with self.subTest(term=term):
                self.assertIn(term, permanent_text)

    def test_missing_optional_permanent_figure_warns_and_still_builds(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            material = deepcopy(generator.PERMANENT)
            missing = Path(temp_dir) / "missing-permanent-figure.png"
            material["figures"][-1] = (missing, material["figures"][-1][1])
            output = Path(temp_dir) / "permanent.docx"
            material["output"] = output

            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter("always")
                self.assertEqual(generator.build_disclosure(material), output)

            self.assertTrue(output.is_file())
            self.assertTrue(
                any(str(missing) in str(item.message) for item in caught),
                [str(item.message) for item in caught],
            )


class DisclosureDocumentIntegrationTests(unittest.TestCase):
    def test_missing_required_strengthened_figures_fail_before_output(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            material = deepcopy(generator.LIGHTWEIGHT)
            missing_one = Path(temp_dir) / "missing-one.png"
            missing_two = Path(temp_dir) / "missing-two.png"
            material["figures"][0] = (missing_one, material["figures"][0][1])
            material["figures"][1] = (missing_two, material["figures"][1][1])
            output = Path(temp_dir) / EXPECTED_OUTPUT_NAME
            material["output"] = output

            with self.assertRaises(FileNotFoundError) as caught:
                generator.build_disclosure(material)

            self.assertIn(str(missing_one), str(caught.exception))
            self.assertIn(str(missing_two), str(caught.exception))
            self.assertFalse(output.exists())

    def test_malformed_template_reports_structure_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            malformed_template = Path(temp_dir) / "malformed.docx"
            Document().save(malformed_template)
            material = deepcopy(generator.LIGHTWEIGHT)
            material["figures"] = []
            material["require_all_figures"] = False
            material["output"] = Path(temp_dir) / EXPECTED_OUTPUT_NAME
            original_template = generator.TEMPLATE
            generator.TEMPLATE = malformed_template
            try:
                with self.assertRaises(Exception) as caught:
                    generator.build_disclosure(material)
            finally:
                generator.TEMPLATE = original_template

            self.assertIs(type(caught.exception), ValueError)
            self.assertIn("至少2个表格", str(caught.exception))
            self.assertFalse(material["output"].exists())

    def test_build_disclosure_writes_strengthened_text_and_exactly_five_images(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            material = deepcopy(generator.LIGHTWEIGHT)
            output = Path(temp_dir) / EXPECTED_OUTPUT_NAME
            material["output"] = output

            self.assertEqual(generator.build_disclosure(material), output)
            self.assertTrue(output.is_file())

            text = docx_text(output)
            for term in [
                EXPECTED_TITLE,
                "CN119312543A",
                "C_gate",
                "R_1m⊆U_A^l×P_+(U_B^l)",
                "DEL=U_A^l-(dom(R)∪REVIEW_A)",
                "d_num(x,y)=clip((|x-y|-eps_abs)/(eps_abs+eps_rel*max(|x|,|y|)),0,1)",
                "I_ms=Σ_t min(c_X(t),c_Y(t))",
                "S(a,b)=Σ alpha_i(1-d_i)/Σ alpha_i",
                "rho_pair=1.2%",
                "94个R_11",
                "DEL=3",
                "ADD=8",
                "图5 场景处理与降级示意图",
            ]:
                with self.subTest(term=term):
                    self.assertIn(term, text)

            with zipfile.ZipFile(output) as package:
                media = {
                    name
                    for name in package.namelist()
                    if name.startswith("word/media/") and not name.endswith("/")
                }
            self.assertEqual(len(media), 5, sorted(media))


if __name__ == "__main__":
    unittest.main()
