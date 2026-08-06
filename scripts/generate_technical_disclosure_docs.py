from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(
    "C:/Users/xiaodi.su/Documents/WXWorkLocalPro_data/Profiles/"
    "32A711E5667E8FD210E3A23C66FA1C18/Cache/chat/file/202607/"
    "专利技术交底书模板_V2_副本.docx"
)
OUT_DIR = ROOT / "docs" / "submission" / "技术交底书"
ASSET_DIR = ROOT / "docs" / "submission" / "assets"


def set_paragraph_text(paragraph, text: str, bold: bool = False, size: int | None = None) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)


def clear_cell(cell) -> None:
    cell._tc.clear_content()
    cell.add_paragraph()


def add_heading(cell, text: str) -> None:
    p = cell.paragraphs[-1] if len(cell.paragraphs) == 1 and not cell.paragraphs[-1].text else cell.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)


def add_para(cell, text: str = "") -> None:
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    if text:
        run = p.add_run(text)
        run.font.size = Pt(10.5)


def add_bullets(cell, items: list[str]) -> None:
    for item in items:
        p = cell.add_paragraph(style=None)
        p.paragraph_format.left_indent = Pt(14)
        p.paragraph_format.first_line_indent = Pt(-10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"• {item}")
        run.font.size = Pt(10.5)


def fill_row(cell, title: str, paragraphs: list[str] | None = None, bullets: list[str] | None = None) -> None:
    clear_cell(cell)
    add_heading(cell, title)
    for text in paragraphs or []:
        add_para(cell, text)
    if bullets:
        add_bullets(cell, bullets)


def add_figures(cell, figures: list[tuple[Path, str]]) -> None:
    for image_path, caption in figures:
        if not image_path.exists():
            continue
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image_path), width=Inches(5.5))
        cap = cell.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        run.font.size = Pt(9)


def fill_metadata(doc: Document, title: str) -> None:
    replacements = {
        2: f"专利申请名称：{title}",
        3: "此专利是否已经公开：否",
        4: "申请人及部门：待补充",
        5: "发明人或撰写人：待补充",
        6: "发明人或撰写人的电话：待补充",
        7: "类型：发明专利",
        8: "内部编号：待补充",
    }
    for idx, text in replacements.items():
        if idx < len(doc.paragraphs):
            set_paragraph_text(doc.paragraphs[idx], text, size=10.5)

    if doc.tables:
        table = doc.tables[0]
        if len(table.rows) >= 3:
            for row_idx in [1, 2]:
                for col_idx in range(1, len(table.rows[row_idx].cells)):
                    table.rows[row_idx].cells[col_idx].text = "待补充"


LIGHTWEIGHT = {
    "short_title": "跨CAD内核模型差异快速对比方法",
    "full_title": "一种面向跨CAD内核的基于轻量化多层摘要的CAD模型差异快速筛选与对比方法及系统",
    "output": OUT_DIR / "跨CAD内核轻量化摘要-CAD模型差异快速筛选与对比-专利技术交底书.docx",
    "figures": [
        (ASSET_DIR / "lightweight-fig1-method.png", "图1 CAD模型差异快速筛选与对比方法流程图"),
        (ASSET_DIR / "lightweight-fig2-system.png", "图2 CAD模型差异快速筛选与对比系统结构图"),
        (ASSET_DIR / "lightweight-fig3-decision.png", "图3 变换判别与多层摘要闸门预筛流程图"),
    ],
    "background": [
        "CAD模型版本比对、供应商模型验收和跨格式转换验证通常需要识别两个三维模型之间的实体级差异和设计语义级差异。现有方法常依赖同一CAD内核下的永久命名、完整BREP逐项比较、关键点匹配、点云距离或图匹配。",
        "在跨CAD内核、跨文件格式或第三方模型交换场景中，永久命名可能缺失、不可比较或不可信；完整BREP结构也可能因内核差异和格式转换而不一致。若直接进行全量BREP实体逐项比较或实体两两几何匹配，计算量容易接近平方级，并且容易把位姿、尺度或仿射变换误判为设计差异。",
    ],
    "solution": [
        "本方案在不要求两个CAD模型存在可比较永久命名、且初始筛选阶段不执行完整BREP实体逐项比较的前提下，先提取全局摘要、局部实体摘要、语义摘要和关系摘要中的至少两类轻量化摘要。",
        "系统基于摘要确定位姿、尺度或仿射变换关系，并进行归一化；对于无法归一化或不满足业务规则的仿射变换，作为全局差异输出。",
        "在执行完整BREP实体逐项比较之前，系统建立摘要索引，并执行全局摘要闸门、局部摘要闸门和关系摘要闸门中的至少两个闸门。每一闸门均包括输入摘要、闸门规则和闸门输出，输出候选关系、排除关系或匹配置信度。",
        "只有通过摘要闸门的候选实体进入局部验证；必要时仅读取或比较候选实体邻域内的局部BREP数据。最终将实体级新增、删除、修改和保持不变结果聚合为孔变化、倒角变化、圆角变化、阵列变化、装配约束变化等设计语义级差异。",
    ],
    "advantages": [
        "适用于跨CAD内核、跨格式转换、供应商模型验收和永久命名不可比较场景。",
        "在完整BREP逐项比较之前筛除大量不相关特征或实体，减少全量几何匹配和BREP读取成本。",
        "通过位姿、尺度和仿射变换处理，降低坐标系变化、单位变化或导入导出变换造成的误判。",
        "输出设计语义级差异，便于工程变更审查、制造影响评估和模型版本管理。",
    ],
    "key_points": [
        "不以两个CAD模型中的面、边或顶点具有相同永久命名为匹配前提。",
        "在完整BREP实体逐项比较之前执行多层轻量化摘要索引和摘要闸门。",
        "摘要闸门的输出是候选关系、排除关系或匹配置信度，而不是最终相同判定。",
        "仅对候选实体或候选邻域执行局部验证，并将实体级差异聚合为设计语义级差异。",
        "不以NURBS控制点关键点匹配、全模型点云距离比较或Ullmann子图同构作为必要步骤。",
    ],
}


PERMANENT = {
    "short_title": "基于永久命名的CAD模型差异对比方法",
    "full_title": "一种基于CAD永久命名的模型差异对比方法及系统",
    "output": OUT_DIR / "永久命名-CAD模型差异对比-专利技术交底书-权利要求增强版.docx",
    "figures": [
        (ASSET_DIR / "permanent_naming-fig1-method.png", "图1 基于永久命名的CAD模型差异对比方法流程图"),
        (ASSET_DIR / "permanent_naming-fig2-system.png", "图2 基于永久命名的CAD模型差异对比系统结构图"),
        (ASSET_DIR / "permanent_naming-fig3-decision.png", "图3 命名映射与命名异常闭环处理流程图"),
    ],
    "background": [
        "同一CAD系统或同一产品数据管理链路中，CAD实体通常具有永久命名标识。直接利用永久命名可以减少全量几何匹配，但永久命名本身并不能说明同名实体是否发生几何、拓扑、参数或语义变化。",
        "实际编辑过程中可能出现命名复用、命名冲突、命名失效、命名漂移、命名分裂或命名合并。若差异对比系统直接信任永久命名，可能将真实修改误判为未变化，或者将命名异常误判为新增/删除。",
        "现有永久命名相关技术更多关注如何生成或维持永久命名，而本方案关注在差异对比流程中如何读取既有命名索引、审计其可信度，并在异常区域通过摘要辅助匹配和局部验证避免差异误判。",
    ],
    "solution": [
        "本方案首先读取由CAD系统、CAD内核或产品数据管理系统预先提供的永久命名实体索引，并判断两个索引是否属于同一命名域，或者是否存在已建立的统一命名域映射关系。",
        "命名域可信度判断：读取命名域标识、CAD系统标识、CAD内核标识、模型版本标识、映射表版本和映射来源，计算永久命名标识覆盖率、永久命名标识唯一性、永久命名冲突率和映射覆盖率等指标。若可信度满足阈值，则进入基于永久命名标识的差异对比流程；若不满足，则拒绝进入该流程，或仅对满足映射置信度的局部实体集合进入流程。",
        "在满足命名域条件后，系统基于永久命名标识执行集合运算，识别新增实体、删除实体和共有实体。集合运算结果作为后续命名一致性审计、版本指纹比较和命名异常闭环处理的基础输入。",
        "命名一致性审计：输入包括永久命名标识、实体类型、父子层级、所属特征、拓扑邻接关系、装配约束关系、版本指纹、实体摘要和映射置信度。审计规则包括：同名实体的摘要或版本指纹差异超过阈值；同名实体的父子层级、所属特征、邻接关系或语义结构发生不满足业务规则的跳转；单侧缺失命名实体但另一模型存在摘要相似且局部关系可信的候选实体；一个命名实体对应多个候选实体或多个命名实体对应一个候选实体组。审计输出为命名正常实体集合、命名异常实体集合、异常类型和置信度。",
        "命名审计评分和异常队列：系统可根据实体摘要一致性、版本指纹一致性、父子层级一致性、拓扑邻接一致性、语义标签一致性、映射置信度和数据来源可信度计算命名审计评分。高于第一阈值的实体进入命名正常集合；介于第一阈值和第二阈值之间的实体进入命名异常集合并执行摘要辅助匹配；低于第二阈值的实体输出为命名异常待复核。命名异常队列记录永久命名标识、异常类型、触发规则、异常前后摘要、替代候选列表、审计评分和局部验证状态。",
        "命名异常闭环：命名异常实体集合中的异常类型至少包括命名复用或冲突、命名失效或疑似重命名、命名漂移、命名分裂和命名合并。系统不直接根据永久命名得出最终差异结论，而是将异常实体转入替代候选流程；根据后续验证结果，将异常实体重分类为新增、删除、修改、重命名、分裂、合并或命名异常待复核。",
        "摘要辅助匹配：对命名异常实体生成替代候选集合。可使用几何摘要、拓扑摘要、参数摘要、语义摘要、关系摘要和外包络摘要中的至少两类摘要；其中摘要用于候选生成和一致性校验，不作为永久命名生成规则，也不作为最终相似判定或最终差异判定的单一依据。候选集合可按实体类型、所属特征、空间邻域、拓扑邻接关系、语义标签和参数摘要检索，并计算候选匹配评分，按评分排序后保留满足条件的替代候选。",
        "局部验证：验证对象包括候选变化实体、命名异常待验证实体以及摘要辅助匹配得到的替代候选集合。验证内容包括局部几何验证、局部拓扑验证、参数验证、语义验证和装配约束验证。对于命名正常且版本指纹一致的共有实体，系统跳过完整几何验证；对于命名异常实体，仅在替代候选范围内执行局部验证。",
        "差异证据记录和增量缓存：系统生成差异证据记录，记录集合运算结果、版本指纹比较结果、命名一致性审计结果、摘要辅助匹配结果、局部验证结果、重分类结果和语义聚合结果。系统还可缓存永久命名实体索引、实体摘要、版本指纹、命名一致性审计结果和命名异常处理结果，以支持增量版本对比。",
        "语义聚合：系统沿实体依赖关系执行影响传播，将实体级新增、删除、修改、重命名、分裂、合并和命名异常待复核结果聚合为孔变化、倒角变化、圆角变化、阵列变化、槽变化、草图变化、参数变化、装配约束变化等设计语义级差异，并输出关联实体、关联特征、关联参数、装配约束、制造特征、检测标注及影响范围。",
    ],
    "advantages": [
        "将全量几何实体匹配转化为命名集合运算、命名一致性审计、指纹筛选和候选局部验证。",
        "不保护永久命名生成机制，避免与已有永久命名生成/维护方案正面重叠。",
        "能够识别命名复用、命名冲突、命名失效、命名漂移、命名分裂和命名合并，降低错误命名导致的差异误判。",
        "版本指纹相同且命名审计正常的共有实体可跳过完整几何验证，提高同一命名域模型版本对比效率。",
        "通过设计语义聚合和影响传播，输出更适合工程审查的差异结果。",
    ],
    "key_points": [
        "永久命名实体索引作为既有输入被读取，不生成或更新CAD系统的永久命名规则。",
        "只有同一命名域或已有统一命名域映射关系、且命名域可信度满足预设条件时，才进入命名集合运算流程。",
        "命名域可信度可由永久命名标识覆盖率、唯一性、冲突率、映射覆盖率和映射置信度等指标确定。",
        "命名一致性审计以永久命名标识、版本指纹、实体摘要、局部拓扑关系、语义结构和依赖关系为输入，输出命名正常实体集合、命名异常实体集合、异常类型和置信度。",
        "命名审计评分用于区分命名正常、命名异常待验证和命名异常待复核实体，异常队列记录触发规则、替代候选、评分和验证状态。",
        "命名异常闭环覆盖命名复用/冲突、命名失效/疑似重命名、命名漂移、命名分裂和命名合并，并将异常实体重分类为新增、删除、修改、重命名、分裂、合并或待复核。",
        "摘要辅助匹配使用几何摘要、拓扑摘要、参数摘要、语义摘要、关系摘要和外包络摘要生成替代候选集合，并可按候选匹配评分排序。",
        "局部验证仅作用于候选变化实体、命名异常待验证实体和替代候选集合；命名正常且版本指纹一致的共有实体跳过完整几何验证。",
        "差异证据记录保存集合运算、审计、匹配、验证、重分类和语义聚合过程，便于审查、追溯和人工复核。",
        "增量缓存保存永久命名索引、实体摘要、版本指纹、审计结果和异常处理结果，以支持后续版本快速对比。",
        "语义聚合将实体级差异、命名异常处理结果和影响传播结果输出为工程可理解的设计语义级差异。",
    ],
}


def build_disclosure(material: dict) -> Path:
    doc = Document(TEMPLATE)
    fill_metadata(doc, material["short_title"])

    detail_table = doc.tables[1]
    fill_row(detail_table.rows[0].cells[0], "专利申请名称", [material["full_title"]])
    fill_row(detail_table.rows[1].cells[0], "同类技术的发展状况、缺陷和不足", material["background"])
    fill_row(detail_table.rows[2].cells[0], "专利技术方案", material["solution"])
    fill_row(
        detail_table.rows[3].cells[0],
        "附图说明",
        [
            "本交底书建议配套以下附图，附图用于说明方法流程、系统模块和关键判断流程。"
        ],
    )
    add_figures(detail_table.rows[3].cells[0], material["figures"])
    fill_row(detail_table.rows[4].cells[0], "本专利技术的优点和有益效果", bullets=material["advantages"])
    fill_row(detail_table.rows[5].cells[0], "需要保护的关键创新点", bullets=material["key_points"])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(material["output"])
    return material["output"]


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    for material in [LIGHTWEIGHT, PERMANENT]:
        print(build_disclosure(material))


if __name__ == "__main__":
    main()
