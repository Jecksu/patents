from __future__ import annotations

import re
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from mathtext import add_math_runs


ROOT = Path(__file__).resolve().parents[1]
PATENT_DIR = ROOT / "docs" / "patent"
OUT_DIR = ROOT / "docs" / "submission"
ASSET_DIR = OUT_DIR / "assets"

# 申请首页信息。代理机构与申请日期留空由代理人填写。
APPLICANT = "苏潇迪-云平台体验研发部"
INVENTOR = "苏潇迪"
INVENTOR_PHONE = "15077207427"


MATERIALS = [
    {
        "key": "lightweight",
        "source": PATENT_DIR / "2026-07-07-cad-model-diff-lightweight-summary-cross-kernel-revised-patent-draft.md",
        "title": "一种面向跨CAD内核的基于轻量化多层摘要的CAD模型差异快速筛选与对比方法及系统",
        "output": OUT_DIR / "跨CAD内核轻量化摘要-CAD模型差异快速筛选与对比-发明专利申请文件.docx",
        "abstract_figure": [
            "获取第一/第二CAD模型",
            "提取多层轻量化摘要",
            "判别位姿/尺度/仿射变换",
            "归一化或记录全局差异",
            "多层摘要闸门粗过滤",
            "候选邻域局部验证",
            "语义聚合",
            "输出差异结果",
        ],
        "system_modules": [
            "模型解析模块",
            "轻量化摘要提取模块",
            "变换判别模块",
            "多层摘要闸门模块",
            "候选匹配模块",
            "局部验证模块",
            "语义差异聚合模块",
            "差异输出模块",
        ],
        "decision_flow": [
            "估计全局变换",
            "刚体变换",
            "相似变换",
            "仿射变换",
            "多层摘要闸门预筛",
            "记录全局形变差异",
        ],
        "figure_titles": [
            "图1 CAD模型差异对比方法流程图",
            "图2 CAD模型差异对比系统结构图",
            "图3 变换判别与归一化流程图",
        ],
    },
    {
        "key": "permanent_naming",
        "source": PATENT_DIR / "2026-07-07-cad-model-diff-permanent-naming-core-patent-draft.md",
        "title": "一种基于CAD永久命名的模型差异对比方法及系统",
        "output": OUT_DIR / "永久命名-CAD模型差异对比-发明专利申请文件.docx",
        "abstract_figure": [
            "读取既有命名索引与关联关系",
            "计算质量状态并三态门控",
            "种子关系验证与三类假设",
            "联合互斥选择与未决冻结",
            "双侧分区检查与失败回流",
            "闭包通过后差异分类",
            "输出类型化事件与证据",
        ],
        "system_modules": [
            "既有命名关系读取模块",
            "质量状态与门控模块",
            "种子验证与假设生成模块",
            "联合选择与冻结模块",
            "分区检查与失败回流模块",
            "闭包后差异分类模块",
            "类型化事件输出模块",
        ],
        "decision_flow": [
            "读取既有命名关系",
            "元数据门控失败停用直配",
            "全局未通过仅逐键可信",
            "全局通过进入种子路径",
            "逐键四状态门控",
            "异常端点转候选假设",
        ],
        "extra_figure": "metrics_diff",
        "figure_titles": [
            "图1 基于永久命名的CAD模型差异对比方法流程图",
            "图2 基于永久命名的CAD模型差异对比系统结构图",
            "图3 命名域元数据硬门控与三态路由流程图",
            "图4 关系闭包与差异分类两支汇合流程图",
        ],
    },
]


def font_path() -> str | None:
    for candidate in [
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]:
        if candidate.exists():
            return str(candidate)
    return None


def get_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    path = font_path()
    if path:
        return ImageFont.truetype(path, size=size, index=1 if bold and path.endswith(".ttc") else 0)
    return ImageFont.load_default()


def wrap_label(label: str, width: int = 12) -> str:
    parts = []
    line = ""
    for ch in label:
        line += ch
        if len(line) >= width:
            parts.append(line)
            line = ""
    if line:
        parts.append(line)
    return "\n".join(parts)


def draw_centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, font) -> None:
    lines = text.splitlines()
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + 5 * (len(lines) - 1)
    y = box[1] + ((box[3] - box[1]) - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        x = box[0] + ((box[2] - box[0]) - w) / 2
        draw.text((x, y), line, fill="black", font=font)
        y += h + 5


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int]) -> None:
    draw.line([start, end], fill="black", width=3)
    x1, y1 = start
    x2, y2 = end
    if y2 > y1:
        points = [(x2, y2), (x2 - 8, y2 - 14), (x2 + 8, y2 - 14)]
    elif y2 < y1:
        points = [(x2, y2), (x2 - 8, y2 + 14), (x2 + 8, y2 + 14)]
    elif x2 > x1:
        points = [(x2, y2), (x2 - 14, y2 - 8), (x2 - 14, y2 + 8)]
    else:
        points = [(x2, y2), (x2 + 14, y2 - 8), (x2 + 14, y2 + 8)]
    draw.polygon(points, fill="black")


def make_vertical_flow(path: Path, title: str, steps: list[str]) -> None:
    w, h = 1500, 1900
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(44, bold=True)
    box_font = get_font(34)
    draw_centered_text(draw, (80, 40, w - 80, 120), title, title_font)
    box_w, box_h = 900, 130
    gap = 52
    x = (w - box_w) // 2
    y = 180
    boxes = []
    for step in steps:
        box = (x, y, x + box_w, y + box_h)
        boxes.append(box)
        draw.rounded_rectangle(box, radius=8, outline="black", width=3, fill="white")
        draw_centered_text(draw, box, wrap_label(step, 13), box_font)
        y += box_h + gap
    for a, b in zip(boxes, boxes[1:]):
        arrow(draw, ((a[0] + a[2]) // 2, a[3]), ((b[0] + b[2]) // 2, b[1]))
    img.save(path)


def make_module_grid(path: Path, title: str, modules: list[str]) -> None:
    w, h = 1700, 1300
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(44, bold=True)
    box_font = get_font(30)
    draw_centered_text(draw, (80, 40, w - 80, 120), title, title_font)
    cols = 3
    box_w, box_h = 460, 140
    gap_x, gap_y = 80, 85
    start_x = (w - cols * box_w - (cols - 1) * gap_x) // 2
    start_y = 210
    for idx, module in enumerate(modules):
        row, col = divmod(idx, cols)
        x = start_x + col * (box_w + gap_x)
        y = start_y + row * (box_h + gap_y)
        box = (x, y, x + box_w, y + box_h)
        draw.rounded_rectangle(box, radius=8, outline="black", width=3, fill="white")
        draw_centered_text(draw, box, wrap_label(module, 10), box_font)
    img.save(path)


def make_decision_flow(path: Path, title: str, labels: list[str]) -> None:
    w, h = 1700, 1200
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(44, bold=True)
    box_font = get_font(30)
    draw_centered_text(draw, (80, 40, w - 80, 120), title, title_font)
    boxes = {
        "top": (620, 180, 1080, 320),
        "left": (160, 470, 550, 610),
        "mid": (655, 470, 1045, 610),
        "right": (1150, 470, 1540, 610),
        "bottom_left": (390, 800, 780, 940),
        "bottom_right": (930, 800, 1320, 940),
    }
    for box, label in zip(boxes.values(), labels):
        draw.rounded_rectangle(box, radius=8, outline="black", width=3, fill="white")
        draw_centered_text(draw, box, wrap_label(label, 10), box_font)
    top_c = ((boxes["top"][0] + boxes["top"][2]) // 2, boxes["top"][3])
    for name in ["left", "mid", "right"]:
        target = ((boxes[name][0] + boxes[name][2]) // 2, boxes[name][1])
        arrow(draw, top_c, target)
    for name in ["left", "mid"]:
        arrow(draw, ((boxes[name][0] + boxes[name][2]) // 2, boxes[name][3]), ((boxes["bottom_left"][0] + boxes["bottom_left"][2]) // 2, boxes["bottom_left"][1]))
    arrow(draw, ((boxes["right"][0] + boxes["right"][2]) // 2, boxes["right"][3]), ((boxes["bottom_right"][0] + boxes["bottom_right"][2]) // 2, boxes["bottom_right"][1]))
    img.save(path)


def make_metrics_diff_flow(path: Path, title: str) -> None:
    """Render the two-branch relation-closure and difference-classification flow."""
    w, h = 1900, 2320
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(42, bold=True)
    box_font = get_font(27)
    small_font = get_font(23)
    draw_centered_text(draw, (70, 30, w - 70, 110), title, title_font)

    left_x, right_x = 90, 1030
    branch_w, box_h = 760, 150
    branch_y = [170, 350, 530, 710, 890]
    trunk_x, trunk_w = 90, 1700
    trunk_y = [1130, 1310, 1490, 1670, 1850, 2030]
    left_labels = [
        "既有永久命名索引与实体—比较规范键多重关系",
        "硬约束／软证据／允许变化属性分离与质量状态",
        "元数据硬门控与全局三态路由、逐键四状态",
        "临时种子关系 R_seed",
        "身份／指纹局部验证",
    ]
    right_labels = [
        "稀疏异常候选边生成与双向索引查询预算",
        "一对一／一对多／多对一三类关系假设",
        "组聚合不变量与局部 B-rep 证据",
        "候选权重与归一化目标 Φ_G",
        "绝对裕量 δ_margin 判定",
    ]
    trunk_labels = [
        "同一局部分量联合端点互斥选择得 R_11／R_1m／R_m1",
        "冻结最小不动点 FROZEN_A／FROZEN_B",
        "end_A／end_B 双侧完备不交分区检查",
        "失败回流：撤销、局部扩读、重建假设、重新选择",
        "CLOSURE_PASS 与 closure_snapshot_id",
        "闭包后分类并物化 CAD_DIFF_EVENT_V1",
    ]

    def draw_column(
        x: int,
        labels: list[str],
        y_values: list[int],
        box_w: int,
        wrap: int,
    ) -> list[tuple[int, int, int, int]]:
        boxes = []
        for y, label in zip(y_values, labels):
            box = (x, y, x + box_w, y + box_h)
            boxes.append(box)
            draw.rounded_rectangle(box, radius=8, outline="black", width=3, fill="white")
            draw_centered_text(draw, box, wrap_label(label, wrap), box_font)
        for a, b in zip(boxes, boxes[1:]):
            arrow(draw, ((a[0] + a[2]) // 2, a[3]), ((b[0] + b[2]) // 2, b[1]))
        return boxes

    left_boxes = draw_column(left_x, left_labels, branch_y, branch_w, 18)
    right_boxes = draw_column(right_x, right_labels, branch_y, branch_w, 18)
    trunk_boxes = draw_column(trunk_x, trunk_labels, trunk_y, trunk_w, 40)

    merge_top = trunk_boxes[0][1]
    left_join = (trunk_x + trunk_w // 3, merge_top)
    right_join = (trunk_x + 2 * trunk_w // 3, merge_top)
    arrow(draw, ((left_boxes[-1][0] + left_boxes[-1][2]) // 2, left_boxes[-1][3]), left_join)
    arrow(draw, ((right_boxes[-1][0] + right_boxes[-1][2]) // 2, right_boxes[-1][3]), right_join)
    draw.text((250, 1070), "门控与种子证据", fill="black", font=small_font)
    draw.text((1450, 1070), "候选与裕量证据", fill="black", font=small_font)
    img.save(path)


def section_between(text: str, start: str, end: str | None = None) -> str:
    start_pat = re.escape(start)
    if end:
        end_pat = re.escape(end)
        m = re.search(start_pat + r"\n(.*?)\n" + end_pat, text, re.S)
    else:
        m = re.search(start_pat + r"\n(.*)", text, re.S)
    return m.group(1).strip() if m else ""


def clean_md_line(line: str) -> str:
    line = line.rstrip()
    line = re.sub(r"^\s*[-*]\s+", "", line)
    line = re.sub(r"^\s*#+\s*", "", line)
    return line


def set_run_font(run, name: str = "SimSun", size: int | None = None, bold: bool | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.1) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def configure_doc(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(11)

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("发明专利申请文件")
    set_run_font(footer_run, size=9)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_title_page(doc: Document, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("发明专利申请文件")
    set_run_font(r, size=22, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run(title)
    set_run_font(r, size=18, bold=True)

    cover_fields = [
        ("申请人", APPLICANT),
        ("发明人", INVENTOR),
        ("联系电话", INVENTOR_PHONE),
        ("代理机构", None),
        ("申请日期", None),
    ]
    for label, value in cover_fields:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(3)
        p.paragraph_format.space_after = Pt(12)
        filler = value if value else "____________________________"
        r = p.add_run(f"{label}：{filler}")
        set_run_font(r, size=12)

    doc.add_page_break()


def add_directory(doc: Document) -> None:
    doc.add_heading("申请文件目录", level=1)
    entries = ["一、说明书摘要", "二、摘要附图", "三、权利要求书", "四、说明书", "五、说明书附图"]
    for entry in entries:
        p = doc.add_paragraph(entry)
        set_paragraph_spacing(p, after=6)
        for run in p.runs:
            set_run_font(run, size=11)
    doc.add_page_break()


def add_md_content(doc: Document, md: str, top_level: int = 1) -> None:
    for raw_line in md.splitlines():
        line = raw_line.rstrip()
        if not line:
            continue
        if line.startswith("#### "):
            p = doc.add_heading(clean_md_line(line), level=min(top_level + 2, 3))
        elif line.startswith("### "):
            p = doc.add_heading(clean_md_line(line), level=min(top_level + 1, 3))
        elif line.startswith("## "):
            p = doc.add_heading(clean_md_line(line), level=top_level)
        else:
            p = doc.add_paragraph()
            text = clean_md_line(line)
            add_math_runs(p, text, lambda run: set_run_font(run, size=11))
            set_paragraph_spacing(p, after=6)
            if re.match(r"^\s+\S", raw_line):
                p.paragraph_format.left_indent = Cm(0.74)
            if re.match(r"^\d+\.\s", text):
                p.paragraph_format.first_line_indent = Cm(-0.35)
                p.paragraph_format.left_indent = Cm(0.35)
        for run in p.runs:
            if p.style.name.startswith("Heading"):
                set_run_font(run, size=int(p.style.font.size.pt), bold=True)


def add_figure(doc: Document, image_path: Path, caption: str, width_in: float = 5.9) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image_path), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run_font(run, size=10)
    set_paragraph_spacing(cap, after=12)


def build_material(material: dict) -> None:
    source_text = material["source"].read_text(encoding="utf-8")
    abstract = section_between(source_text, "## 摘要", "## 权利要求书")
    claims = section_between(source_text, "## 权利要求书", "## 说明书")
    specification = section_between(source_text, "## 说明书")

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    fig1 = ASSET_DIR / f"{material['key']}-fig1-method.png"
    fig2 = ASSET_DIR / f"{material['key']}-fig2-system.png"
    fig3 = ASSET_DIR / f"{material['key']}-fig3-decision.png"
    make_vertical_flow(fig1, material["figure_titles"][0], material["abstract_figure"])
    make_module_grid(fig2, material["figure_titles"][1], material["system_modules"])
    make_decision_flow(fig3, material["figure_titles"][2], material["decision_flow"])
    extra_fig = None
    if material.get("extra_figure") == "metrics_diff":
        extra_fig = ASSET_DIR / f"{material['key']}-fig4-metrics-diff.png"
        make_metrics_diff_flow(extra_fig, material["figure_titles"][3])

    doc = Document()
    configure_doc(doc, material["title"])
    add_title_page(doc, material["title"])
    add_directory(doc)

    doc.add_heading("一、说明书摘要", level=1)
    add_md_content(doc, abstract, top_level=2)
    doc.add_page_break()

    doc.add_heading("二、摘要附图", level=1)
    add_figure(doc, fig1, "摘要附图")
    doc.add_page_break()

    doc.add_heading("三、权利要求书", level=1)
    add_md_content(doc, claims, top_level=2)
    doc.add_page_break()

    doc.add_heading("四、说明书", level=1)
    add_md_content(doc, specification, top_level=2)
    doc.add_page_break()

    doc.add_heading("五、说明书附图", level=1)
    figures = [fig1, fig2, fig3] + ([extra_fig] if extra_fig else [])
    for fig, caption in zip(figures, material["figure_titles"]):
        add_figure(doc, fig, caption)

    material["output"].parent.mkdir(parents=True, exist_ok=True)
    doc.save(material["output"])


def main() -> None:
    for material in MATERIALS:
        build_material(material)
        print(material["output"])


if __name__ == "__main__":
    main()
